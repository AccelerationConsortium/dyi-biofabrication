#!/usr/bin/env python3
"""Extract itemized bills of materials from supplementary files.

Most papers report only a total or category-level cost; the line-by-line BOM
usually lives in an external repository. This pulls the itemized BOMs out of the
supplementary files we actually hold and writes lists/derived/bom_items.csv.
"""
import csv
import os
import re
import zipfile

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SUPP = os.path.join(REPO, "pdfs", "supplementary")
OUT = os.path.join(REPO, "lists", "derived", "bom_items.csv")

FIELDS = ["paper", "source_file", "item", "notes", "quantity",
          "unit_cost", "line_price", "currency", "supplier", "part_number", "link"]


def clean(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


def from_openworkstation(path):
    """bom_split sheet: module | designator | component | supplier | part-number | n | length | unit cost | price USD"""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    if "bom_split" not in wb.sheetnames:
        return []
    ws = wb["bom_split"]
    rows = list(ws.iter_rows(values_only=True))
    header = [clean(c).lower() for c in rows[0]]

    def col(*names):
        for i, h in enumerate(header):
            if any(n in h for n in names):
                return i
        return None

    i_comp, i_sup = col("component"), col("manufacturer", "supplier")
    i_part, i_qty = col("part-number", "part number"), col("number of item")
    i_unit, i_price = col("cost per unit", "cost per item"), col("price")
    out = []
    for r in rows[1:]:
        comp = clean(r[i_comp]) if i_comp is not None and i_comp < len(r) else ""
        price = clean(r[i_price]) if i_price is not None and i_price < len(r) else ""
        if not comp or comp.lower() in ("component", "total"):
            continue
        out.append({
            "item": comp,
            "notes": clean(r[0]) if r else "",
            "quantity": clean(r[i_qty]) if i_qty is not None and i_qty < len(r) else "",
            "unit_cost": clean(r[i_unit]) if i_unit is not None and i_unit < len(r) else "",
            "line_price": price,
            "currency": "USD",
            "supplier": clean(r[i_sup]) if i_sup is not None and i_sup < len(r) else "",
            "part_number": clean(r[i_part]) if i_part is not None and i_part < len(r) else "",
            "link": "",
        })
    return out


def from_openlh(zip_path):
    """BOM.pdf inside the OpenLH archive: Item Number | Item Name | Notes | Quantity | Price | Purchase Link"""
    from pypdf import PdfReader
    tmp = "/tmp/_bom_openlh.pdf"
    with zipfile.ZipFile(zip_path) as z:
        names = [n for n in z.namelist() if n.lower().endswith("bom.pdf")]
        if not names:
            return []
        with open(tmp, "wb") as f:
            f.write(z.read(names[0]))
    text = " ".join(" ".join(p.extract_text().split()) for p in PdfReader(tmp).pages)
    text = re.sub(r"^.*?Purchase Link", "", text, count=1)
    # rows start with an item number and run to the next item number
    parts = re.split(r"\s(?=\d{1,2}\s+[A-Z3])", text)
    out = []
    for p in parts:
        m = re.match(r"(\d{1,2})\s+(.*)", p.strip())
        if not m:
            continue
        body = m.group(2)
        link = ""
        lm = re.search(r"https?://\S+", body)
        if lm:
            link = lm.group(0).rstrip(".,")
            body = body[:lm.start()].strip()
        price = ""
        pm = re.search(r"(\d[\d,.]*)\s*\$", body)
        if pm:
            price = pm.group(1)
            body = (body[:pm.start()] + body[pm.end():]).strip()
        qty = ""
        qm = re.search(r"\s(\d{1,3})\s*$", body)
        if qm:
            qty = qm.group(1)
            body = body[:qm.start()].strip()
        if not body:
            continue
        out.append({"item": body[:90], "notes": "", "quantity": qty, "unit_cost": "",
                    "line_price": price, "currency": "USD", "supplier": "",
                    "part_number": "", "link": link})
    return out


def main():
    sources = []
    for f in sorted(os.listdir(SUPP)):
        p = os.path.join(SUPP, f)
        paper = re.sub(r"^\d+_", "", f).replace("_supplementary", "").rsplit(".", 1)[0].replace("_", " ")
        if f.endswith(".xlsx") and "OpenWorkstation" in f:
            sources.append((paper, f, from_openworkstation(p)))
        elif f.endswith(".zip") and "OpenLH" in f:
            sources.append((paper, f, from_openlh(p)))

    rows = []
    for paper, fname, items in sources:
        for it in items:
            it["paper"] = paper
            it["source_file"] = fname
            rows.append({k: it.get(k, "") for k in FIELDS})

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(rows)

    print(f"wrote {len(rows)} BOM line items -> {OUT}\n")
    for paper, fname, items in sources:
        priced = sum(1 for i in items if i["line_price"])
        linked = sum(1 for i in items if i["link"])
        print(f"  {paper[:44]:46s} {len(items):3d} items, {priced:3d} priced, {linked:3d} with purchase link")


if __name__ == "__main__":
    main()
